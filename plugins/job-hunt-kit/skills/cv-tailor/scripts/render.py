#!/usr/bin/env python3
"""Render an application's content.json into LaTeX sources and/or Word documents.

content.json is the single source of tailored content for one application (see
references/content-schema.md). This script turns it into:

  latex : cv.tex + cover-letter.tex (+ a copy of jobcv.sty if the folder has none)
  docx  : cv.docx + cover-letter.docx, styled from a Word template

It only writes sources. build.py calls it, then compiles / converts and checks pages.

Usage:
    python render.py FOLDER --format latex|docx|both [--templates DIR]

--templates is the workspace's override folder: DIR/jobcv.sty and DIR/jobcv-template.docx
are used instead of the bundled ones when they exist.

LaTeX rendering is standard library only. Word rendering needs python-docx
(pip install python-docx).
"""

from __future__ import annotations

import argparse
import json
import re
import shutil
import sys
from pathlib import Path

PLUGIN_ROOT = Path(__file__).resolve().parents[3]
BUNDLED_STY = PLUGIN_ROOT / "latex" / "jobcv.sty"
BUNDLED_DOCX = PLUGIN_ROOT / "word" / "jobcv-template.docx"


class ContentError(Exception):
    pass


# =====================================================================
#  Shared: loading, validation, eligibility line, inline markup
# =====================================================================

def load_content(folder: Path) -> dict:
    path = folder / "content.json"
    if not path.exists():
        raise ContentError(f"{path} not found")
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except json.JSONDecodeError as e:
        raise ContentError(f"content.json is not valid JSON: {e}") from e
    for key in ("identity", "cv"):
        if key not in data:
            raise ContentError(f"content.json is missing '{key}'")
    for key in ("name", "email"):
        if not data["identity"].get(key):
            raise ContentError(f"identity.{key} is empty")
    if not data["cv"].get("title"):
        raise ContentError("cv.title is empty (use the exact job title from the posting)")
    return data


def resolve_eligibility(data: dict) -> dict:
    """Validate the eligibility block and normalise the country's spelling.

    The posting's country is REQUIRED: a hand-typed eligibility line is right for one
    country and plausibly wrong for every other, so a missing country is an error rather
    than a guess. If the country matches a listed country ignoring case/spaces, it is
    rewritten to the listed spelling, so LaTeX's exact-text comparison agrees with ours.
    """
    el = dict(data.get("eligibility") or {})
    country = (el.get("country") or "").strip()
    if not country:
        raise ContentError("eligibility.country is not set: it must be the country of THIS posting")
    eligible = [c.strip() for c in (el.get("eligible_countries") or []) if c and c.strip()]
    key = lambda s: re.sub(r"\s+", " ", s).casefold()
    for c in eligible:
        if key(c) == key(country):
            country = c
            break
    el["country"] = country
    el["eligible_countries"] = eligible
    el["is_eligible"] = country in eligible
    return el


def permit_text(el: dict) -> str:
    cit = (el.get("citizenship") or "").strip()
    note = (el.get("permit_note") or "").strip()
    if el["is_eligible"]:
        core = f"eligible to work in {el['country']}"
    else:
        core = f"open to relocation to {el['country']}" + (f" ({note})" if note else "")
    return f"{cit} citizen, {core}" if cit else core[0].upper() + core[1:]


TOKEN_RE = re.compile(r"(\*\*.+?\*\*|\[[^\]]+\]\([^)\s]+\))")
LINK_RE = re.compile(r"\[([^\]]+)\]\(([^)\s]+)\)")


def smart_quotes(s: str) -> str:
    return re.sub(r'"([^"]*)"', "“\\1”", s)


def tokens(text: str):
    """Yield (kind, text, url) for the tiny inline markup: **bold** and [label](url)."""
    for part in TOKEN_RE.split(smart_quotes(text)):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            yield "bold", part[2:-2], None
        else:
            m = LINK_RE.fullmatch(part)
            if m:
                yield "link", m.group(1), m.group(2)
            else:
                yield "text", part, None


def as_list(v) -> list:
    if v is None:
        return []
    return v if isinstance(v, list) else [v]


# =====================================================================
#  LaTeX
# =====================================================================

TEX_ESC = {
    "\\": r"\textbackslash{}", "&": r"\&", "%": r"\%", "$": r"\$", "#": r"\#",
    "_": r"\_", "{": r"\{", "}": r"\}", "~": r"\textasciitilde{}",
    "^": r"\textasciicircum{}", " ": "~",
}


def tex_esc(s: str) -> str:
    return "".join(TEX_ESC.get(ch, ch) for ch in s)


def tex_url(u: str) -> str:
    return u.replace("\\", "/").replace("%", r"\%").replace("#", r"\#")


def tex_inline(text: str) -> str:
    out = []
    for kind, t, url in tokens(text):
        if kind == "bold":
            out.append(r"{\bfseries " + tex_esc(t) + "}")
        elif kind == "link":
            out.append(r"\weblink{" + tex_url(url) + "}{" + tex_esc(t) + "}")
        else:
            out.append(tex_esc(t))
    return "".join(out)


def tex_identity(data: dict, title: str, el: dict | None) -> list[str]:
    idn = data["identity"]
    L = [
        r"\cvname{" + tex_esc(idn.get("name_with_title") or idn["name"]) + "}",
        r"\cvtitle{" + tex_esc(title) + "}",
        r"\cvlocation{" + tex_esc(idn.get("location", "")) + "}",
        r"\cvphone{" + tex_esc(idn.get("phone", "")) + "}",
        r"\cvemail{" + tex_esc(idn["email"]) + "}",
    ]
    for key, macro in (("website", "cvsite"), ("linkedin", "cvlinkedin")):
        v = idn.get(key) or {}
        if v.get("label") and v.get("url"):
            L.append(f"\\{macro}{{{tex_esc(v['label'])}}}{{{tex_url(v['url'])}}}")
    if idn.get("photo"):
        L.append(r"\cvphoto{" + Path(idn["photo"]).name + "}")
    if el is not None:
        if el.get("citizenship"):
            L.append(r"\cvcitizenship{" + tex_esc(el["citizenship"]) + "}")
        L.append(r"\cveligiblein{" + tex_esc(", ".join(el["eligible_countries"])) + "}")
        L.append(r"\cvcountry{" + tex_esc(el["country"]) + "}          % country of THIS posting")
        if el.get("permit_note"):
            L.append(r"\cvpermitnote{" + tex_esc(el["permit_note"]) + "}")
    return L


def tex_quality_block(items: list[str], cols: int, extra: list[str] = ()) -> list[str]:
    body = [r"\quality{" + tex_inline(i) + "}" for i in items] + list(extra)
    if cols <= 1:
        return body
    return [r"\compactcols{" + str(cols) + "}{%"] + body + ["}"]


def tex_section(sec: dict) -> list[str]:
    t = sec.get("type")
    L = [r"\cvsection{" + tex_esc(sec.get("heading", "")) + "}"]
    if t == "paragraph":
        L.append("\n\n".join(tex_inline(p) for p in as_list(sec.get("text"))))
    elif t in ("qualities", "glance"):
        extra = [r"\quality{\permitline}"] if t == "glance" else []
        L += tex_quality_block(as_list(sec.get("items")), int(sec.get("columns", 3)), extra)
    elif t == "skills":
        groups = [r"\skillgroup{" + tex_esc(g["label"]) + "}{" + tex_inline(g["text"]) + "}"
                  for g in sec.get("groups", [])]
        L.append("\n\n".join(groups))
    elif t == "experience":
        for e in sec.get("entries", []):
            L.append(r"\cventry{" + tex_esc(e["role"]) + "}{" + tex_esc(e.get("org", "")) + "}{"
                     + tex_esc(e.get("dates", "")) + "}")
            if e.get("bullets"):
                L.append(r"\begin{cvbullets}")
                L += ["  \\item " + tex_inline(b) for b in e["bullets"]]
                L.append(r"\end{cvbullets}")
            L.append("")
    elif t == "education":
        for e in sec.get("entries", []):
            L.append(r"\eduentrywide{" + tex_esc(e["degree"]) + "}{" + tex_esc(e.get("institution", ""))
                     + "}{" + tex_esc(e.get("dates", "")) + "}")
    elif t == "languages":
        parts = [r"{\bfseries " + tex_esc(i["language"]) + "} --- " + tex_esc(i.get("level", ""))
                 for i in sec.get("items", [])]
        L.append(r"\langline{" + " \\textperiodcentered\\ ".join(parts) + "}")
    elif t == "list":
        L.append(r"\begin{cvbullets}")
        L += ["  \\item " + tex_inline(b) for b in sec.get("items", [])]
        L.append(r"\end{cvbullets}")
    else:
        raise ContentError(f"unknown section type: {t!r}")
    L.append("")
    return L


def render_latex(data: dict, folder: Path, templates: Path | None) -> list[Path]:
    paper = data.get("paper", "a4paper")
    el = resolve_eligibility(data)
    cv = data["cv"]
    written = []

    head = [
        "% Generated from content.json by job-hunt-kit render.py. Edit content.json and",
        "% rebuild; or edit this file and build with --no-render.",
        f"\\documentclass[12pt,{paper}]{{article}}",
        r"\usepackage{jobcv}",
    ]
    hyph = cv.get("hyphenation_exceptions") or []
    if hyph:
        head.append(r"\hyphenation{" + " ".join(hyph) + "}")

    has_glance = any(s.get("type") == "glance" for s in cv.get("sections", []))
    if not has_glance:
        raise ContentError("the CV needs a 'glance' section: it carries the work-eligibility line")

    L = head + [""] + tex_identity(data, cv["title"], el) + [
        "", r"\begin{document}", r"\pagestyle{cvrest}\thispagestyle{cvfirst}", r"\makecvheader", ""]
    for sec in cv.get("sections", []):
        L += tex_section(sec)
    L.append(r"\end{document}")
    (folder / "cv.tex").write_text("\n".join(L) + "\n", encoding="utf-8")
    written.append(folder / "cv.tex")

    letter = data.get("letter")
    if letter:
        title = letter.get("title") or cv["title"]
        L = head + [""] + tex_identity(data, title, None) + [
            "", r"\begin{document}", r"\pagestyle{empty}", r"\makecvheader", ""]
        if letter.get("date"):
            L.append(r"\cldate{" + tex_esc(letter["date"]) + "}")
        rec = as_list(letter.get("recipient"))
        if rec:
            L.append(r"\clrecipient{%" + "\n" + "\\\\\n".join(tex_esc(r) for r in rec) + "}")
        L.append(r"\clsubject{" + tex_esc(letter.get("subject", f"Application: {title}")) + "}")
        L.append(r"\clsalutation{" + tex_esc(letter.get("salutation", "Dear Hiring Team,")) + "}")
        L.append(r"\begin{clbody}")
        L.append("\n\n".join(tex_inline(p) for p in letter.get("paragraphs", [])))
        L.append(r"\end{clbody}")
        L.append(r"\clclosing{" + tex_esc(letter.get("sign_off", "Kind regards,")) + "}{"
                 + tex_esc(letter.get("signature_name") or data["identity"]["name"]) + "}")
        L.append(r"\end{document}")
        (folder / "cover-letter.tex").write_text("\n".join(L) + "\n", encoding="utf-8")
        written.append(folder / "cover-letter.tex")

    # Copy the style in (once) so the folder stays rebuildable if the template changes later.
    sty_dst = folder / "jobcv.sty"
    if not sty_dst.exists():
        src = templates / "jobcv.sty" if templates and (templates / "jobcv.sty").exists() else BUNDLED_STY
        shutil.copy2(src, sty_dst)
    photo = data["identity"].get("photo")
    if photo and not (folder / Path(photo).name).exists():
        raise ContentError(f"photo {Path(photo).name} not found in the application folder: copy it in")
    return written


# =====================================================================
#  Word
# =====================================================================

NAVY = "0D2D5A"
LINKBLUE = "2A5699"

# name: (base, size_pt, bold, italic, colour, all_caps, before_pt, after_pt, keep_next, align, border)
PARA_STYLES = {
    "CV Name":        (None, 20, True, False, NAVY, False, 0, 0, True, None, None),
    "CV Title":       (None, 11.5, False, False, NAVY, True, 2, 2, True, None, None),
    "CV Contact":     (None, 9, False, False, None, False, 0, 10, False, None, ("bottom", 18, NAVY, 6)),
    "CV Section":     (None, 12.5, True, False, NAVY, True, 10, 5, True, None, ("bottom", 8, NAVY, 2)),
    "CV Body":        (None, None, False, False, None, False, 0, 0, False, None, None),
    "CV Role":        (None, 12, True, False, None, False, 7, 0, True, None, None),
    "CV Org":         (None, 11, False, False, NAVY, False, 1, 2, True, None, None),
    "CV Bullet":      ("List Bullet", None, False, False, None, False, 1, 1, False, None, None),
    "CV Skill":       (None, None, False, False, None, False, 3, 0, False, None, None),
    "CV Inline":      (None, None, False, False, None, False, 2, 0, False, None, None),
    "CV Degree":      (None, None, True, False, None, False, 5, 0, True, None, None),
    "CV Institution": (None, 9.5, False, False, NAVY, False, 0, 0, False, None, None),
    "CV Running":     (None, 9, False, False, NAVY, False, 0, 0, False, None, None),
    "CL Date":        (None, 10.5, False, False, NAVY, False, 0, 10, False, None, None),
    "CL Recipient":   (None, 11, False, False, None, False, 0, 0, False, None, None),
    "CL Subject":     (None, 14, True, False, NAVY, False, 14, 10, True, None, ("bottom", 8, NAVY, 3)),
    "CL Salutation":  (None, None, False, False, None, False, 0, 6, True, None, None),
    "CL Body":        (None, None, False, False, None, False, 0, 8, False, "justify", None),
    "CL Closing":     (None, None, False, False, None, False, 10, 10, True, None, None),
    "CL Name":        (None, None, True, False, NAVY, False, 0, 2, True, None, None),
    "CL Signature":   (None, 10, False, False, None, False, 0, 0, False, None, None),
}
CHAR_STYLES = {
    "CV Label": (True, False, NAVY),
    "CV Dates": (False, True, None),
    "CV Link":  (False, False, LINKBLUE),
}


def _docx():
    try:
        import docx  # noqa: F401
    except ImportError as e:
        raise ContentError("python-docx is not installed: run  pip install python-docx") from e


def _border(pPr, side, size_eighths, color, space_pt):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    bdr = pPr.find(qn("w:pBdr"))
    if bdr is None:
        bdr = OxmlElement("w:pBdr")
        pPr.append(bdr)
    el = OxmlElement(f"w:{side}")
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), str(size_eighths))
    el.set(qn("w:space"), str(space_pt))
    el.set(qn("w:color"), color)
    bdr.append(el)


def ensure_styles(doc, font="Calibri", body_pt=11.5):
    """Create the CV/CL styles if the template lacks them. Existing styles win, so a
    template restyled in Word keeps its look."""
    from docx.enum.style import WD_STYLE_TYPE
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor

    styles = doc.styles
    normal = styles["Normal"]
    if "CV Body" not in styles:  # fresh template: set the document defaults too
        normal.font.name = font
        normal.font.size = Pt(body_pt)
        normal.paragraph_format.line_spacing = 1.15
        normal.paragraph_format.space_after = Pt(0)
        from docx.oxml.ns import qn
        rfonts = normal.element.get_or_add_rPr().find(qn("w:rFonts"))
        if rfonts is not None:  # also cover East Asian / complex-script runs
            rfonts.set(qn("w:eastAsia"), font)
            rfonts.set(qn("w:cs"), font)

    for name, (base, size, bold, italic, color, caps, before, after, keep, align, border) in PARA_STYLES.items():
        if name in styles:
            continue
        st = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        st.base_style = styles[base] if base and base in styles else normal
        st.quick_style = True
        f = st.font
        if size:
            f.size = Pt(size)
        f.bold = bold or None
        f.italic = italic or None
        if color:
            f.color.rgb = RGBColor.from_string(color)
        if caps:
            f.all_caps = True
        pf = st.paragraph_format
        pf.space_before = Pt(before)
        pf.space_after = Pt(after)
        pf.keep_with_next = keep or None
        if align == "justify":
            pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if border:
            _border(st.element.get_or_add_pPr(), *border)
        if name == "CV Bullet" and (not base or base not in styles):
            pf.left_indent = Pt(12)
            pf.first_line_indent = Pt(-9)

    for name, (bold, italic, color) in CHAR_STYLES.items():
        if name in styles:
            continue
        st = styles.add_style(name, WD_STYLE_TYPE.CHARACTER)
        st.font.bold = bold or None
        st.font.italic = italic or None
        if color:
            st.font.color.rgb = RGBColor.from_string(color)


def _clear_body(doc):
    from docx.oxml.ns import qn
    body = doc.element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def _hyperlink(par, url, text, char_style=None):
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    r_id = par.part.relate_to(url, RT.HYPERLINK, is_external=True)
    h = OxmlElement("w:hyperlink")
    h.set(qn("r:id"), r_id)
    r = OxmlElement("w:r")
    if char_style:
        rpr = OxmlElement("w:rPr")
        rs = OxmlElement("w:rStyle")
        rs.set(qn("w:val"), par.part.document.styles[char_style].style_id)
        rpr.append(rs)
        r.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    r.append(t)
    h.append(r)
    par._p.append(h)


def _inline(par, text):
    for kind, t, url in tokens(text):
        if kind == "bold":
            par.add_run(t).bold = True
        elif kind == "link":
            _hyperlink(par, url, t, "CV Link")
        else:
            par.add_run(t)


def _field(par, instr):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), instr)
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "1"
    r.append(t)
    fld.append(r)
    par._p.append(fld)


def _text_width(section):
    return section.page_width - section.left_margin - section.right_margin


def _right_tab(par, section):
    from docx.enum.text import WD_TAB_ALIGNMENT
    par.paragraph_format.tab_stops.add_tab_stop(_text_width(section), WD_TAB_ALIGNMENT.RIGHT)


def _setup_page(doc, paper):
    from docx.shared import Mm
    s = doc.sections[0]
    if paper == "letterpaper":
        s.page_width, s.page_height = Mm(215.9), Mm(279.4)
    else:
        s.page_width, s.page_height = Mm(210), Mm(297)
    return s


def _header_block(doc, data, title):
    """Name, title and contact line in the BODY (not the page header, which some ATS
    skip). Plain text with middle-dot separators; links are real hyperlinks."""
    idn = data["identity"]
    p = doc.add_paragraph(style="CV Name")
    p.add_run(idn.get("name_with_title") or idn["name"])
    if idn.get("photo"):
        _anchor_photo(doc, p, idn["photo"])
    doc.add_paragraph(title, style="CV Title")
    c = doc.add_paragraph(style="CV Contact")
    parts = [("t", idn.get("location", "")), ("t", idn.get("phone", "")),
             ("l", (idn["email"], "mailto:" + idn["email"]))]
    for key in ("website", "linkedin"):
        v = idn.get(key) or {}
        if v.get("label") and v.get("url"):
            parts.append(("l", (v["label"], v["url"])))
    first = True
    for kind, val in parts:
        if not val:
            continue
        if not first:
            c.add_run("  ·  ")
        first = False
        if kind == "t":
            c.add_run(val)
        else:
            _hyperlink(c, val[1], val[0])


def _anchor_photo(doc, par, photo):
    """Float a square photo at the top right of page 1 (anchored, not inline, so it
    doesn't push the text down)."""
    from docx.shared import Mm
    from docx.oxml import parse_xml
    path = Path(photo)
    if not path.exists():
        raise ContentError(f"photo {photo} not found")
    run = par.add_run()
    inline = run.add_picture(str(path), width=Mm(24), height=Mm(24))._inline
    extent = inline.find("{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}extent")
    cx, cy = extent.get("cx"), extent.get("cy")
    graphic = inline.find("{http://schemas.openxmlformats.org/drawingml/2006/main}graphic")
    docpr = inline.find("{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}docPr")
    anchor = parse_xml(
        '<wp:anchor xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
        'distT="0" distB="0" distL="114300" distR="0" simplePos="0" relativeHeight="1" '
        'behindDoc="0" locked="0" layoutInCell="1" allowOverlap="1">'
        '<wp:simplePos x="0" y="0"/>'
        '<wp:positionH relativeFrom="margin"><wp:align>right</wp:align></wp:positionH>'
        '<wp:positionV relativeFrom="margin"><wp:posOffset>0</wp:posOffset></wp:positionV>'
        f'<wp:extent cx="{cx}" cy="{cy}"/><wp:effectExtent l="0" t="0" r="0" b="0"/>'
        '<wp:wrapSquare wrapText="left"/></wp:anchor>')
    anchor.append(docpr)
    anchor.append(parse_xml('<wp:cNvGraphicFramePr xmlns:wp="http://schemas.openxmlformats.org/'
                            'drawingml/2006/wordprocessingDrawing"/>'))
    anchor.append(graphic)
    inline.getparent().replace(inline, anchor)


def _running_head(doc, section, data, title, with_head: bool):
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    section.different_first_page_header_footer = True
    for hdr in (section.header, section.first_page_header):
        for p in hdr.paragraphs:
            p.clear()
    if with_head:
        hp = section.header.paragraphs[0]
        hp.style = doc.styles["CV Running"]
        _right_tab(hp, section)
        hp.add_run(data["identity"].get("name_with_title") or data["identity"]["name"]).bold = True
        hp.add_run("\t" + title)
        _border(hp._p.get_or_add_pPr(), "bottom", 4, NAVY, 2)
    for ftr in (section.footer, section.first_page_footer):
        fp = ftr.paragraphs[0]
        fp.clear()
        fp.style = doc.styles["CV Running"]
        fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        if with_head:
            _field(fp, "PAGE")
            fp.add_run(" / ")
            _field(fp, "NUMPAGES")


def _open_template(templates):
    from docx import Document
    tpl = templates / "jobcv-template.docx" if templates and (templates / "jobcv-template.docx").exists() else BUNDLED_DOCX
    doc = Document(str(tpl)) if tpl.exists() else Document()
    ensure_styles(doc)
    _clear_body(doc)
    return doc


def _meta(doc, data, title):
    cp = doc.core_properties
    cp.author = data["identity"]["name"]
    cp.title = f"{data['identity']['name']} - {title}"
    cp.subject = title
    cp.comments = ""
    cp.last_modified_by = data["identity"]["name"]


def build_cv_docx(doc, data, el):
    cv = data["cv"]
    s = _setup_page(doc, data.get("paper", "a4paper"))
    _running_head(doc, s, data, cv["title"], with_head=True)
    _header_block(doc, data, cv["title"])
    for sec in cv.get("sections", []):
        t = sec.get("type")
        doc.add_paragraph(sec.get("heading", ""), style="CV Section")
        if t == "paragraph":
            for ptxt in as_list(sec.get("text")):
                _inline(doc.add_paragraph(style="CV Body"), ptxt)
        elif t in ("qualities", "glance"):
            items = list(as_list(sec.get("items")))
            p = doc.add_paragraph(style="CV Inline")
            for i, it in enumerate(items):
                if i:
                    p.add_run("  ·  ")
                _inline(p, it)
            if t == "glance":
                if items:
                    p.add_run("  ·  ")
                p.add_run(permit_text(el))
        elif t == "skills":
            for g in sec.get("groups", []):
                p = doc.add_paragraph(style="CV Skill")
                p.add_run(g["label"], style="CV Label")
                p.add_run(" — ")
                _inline(p, g["text"])
        elif t == "experience":
            for e in sec.get("entries", []):
                doc.add_paragraph(e["role"], style="CV Role")
                p = doc.add_paragraph(style="CV Org")
                _right_tab(p, s)
                p.add_run(e.get("org", ""))
                if e.get("dates"):
                    p.add_run("\t")
                    p.add_run(e["dates"], style="CV Dates")
                for b in e.get("bullets", []):
                    _inline(doc.add_paragraph(style="CV Bullet"), b)
        elif t == "education":
            for e in sec.get("entries", []):
                p = doc.add_paragraph(style="CV Degree")
                _right_tab(p, s)
                p.add_run(e["degree"])
                if e.get("dates"):
                    p.add_run("\t")
                    r = p.add_run(e["dates"], style="CV Dates")
                    r.bold = False
                doc.add_paragraph(e.get("institution", ""), style="CV Institution")
        elif t == "languages":
            p = doc.add_paragraph(style="CV Inline")
            for i, it in enumerate(sec.get("items", [])):
                if i:
                    p.add_run("  ·  ")
                p.add_run(it["language"]).bold = True
                p.add_run(" — " + it.get("level", ""))
        elif t == "list":
            for b in sec.get("items", []):
                _inline(doc.add_paragraph(style="CV Bullet"), b)
        else:
            raise ContentError(f"unknown section type: {t!r}")
    _meta(doc, data, cv["title"])


def build_letter_docx(doc, data):
    letter = data["letter"]
    title = letter.get("title") or data["cv"]["title"]
    s = _setup_page(doc, data.get("paper", "a4paper"))
    _running_head(doc, s, data, title, with_head=False)
    _header_block(doc, data, title)
    if letter.get("date"):
        doc.add_paragraph(letter["date"], style="CL Date")
    rec = as_list(letter.get("recipient"))
    for i, line in enumerate(rec):
        p = doc.add_paragraph(line, style="CL Recipient")
    doc.add_paragraph(smart_quotes(letter.get("subject", f"Application: {title}")), style="CL Subject")
    doc.add_paragraph(letter.get("salutation", "Dear Hiring Team,"), style="CL Salutation")
    for ptxt in letter.get("paragraphs", []):
        _inline(doc.add_paragraph(style="CL Body"), ptxt)
    doc.add_paragraph(letter.get("sign_off", "Kind regards,"), style="CL Closing")
    doc.add_paragraph(letter.get("signature_name") or data["identity"]["name"], style="CL Name")
    idn = data["identity"]
    sig = doc.add_paragraph(style="CL Signature")
    site = idn.get("website") or {}
    if site.get("label") and site.get("url"):
        _hyperlink(sig, site["url"], site["label"], "CV Link")
        sig.add_run("  ·  ")
    _hyperlink(sig, "mailto:" + idn["email"], idn["email"], "CV Link")
    if idn.get("phone"):
        sig.add_run("  ·  " + idn["phone"])
    _meta(doc, data, title)


def render_docx(data: dict, folder: Path, templates: Path | None) -> list[Path]:
    _docx()
    el = resolve_eligibility(data)
    if not any(s.get("type") == "glance" for s in data["cv"].get("sections", [])):
        raise ContentError("the CV needs a 'glance' section: it carries the work-eligibility line")
    photo = data["identity"].get("photo")
    if photo:
        data = json.loads(json.dumps(data))
        data["identity"]["photo"] = str(folder / Path(photo).name)
    written = []
    doc = _open_template(templates)
    build_cv_docx(doc, data, el)
    doc.save(str(folder / "cv.docx"))
    written.append(folder / "cv.docx")
    if data.get("letter"):
        doc = _open_template(templates)
        build_letter_docx(doc, data)
        doc.save(str(folder / "cover-letter.docx"))
        written.append(folder / "cover-letter.docx")
    return written


# =====================================================================

def render(folder: Path, fmt: str, templates: Path | None = None) -> list[Path]:
    data = load_content(folder)
    out = []
    if fmt in ("latex", "both"):
        out += render_latex(data, folder, templates)
    if fmt in ("docx", "both"):
        out += render_docx(data, folder, templates)
    return out


def make_template(path: Path) -> None:
    """Write the bundled Word template: all styles, plus a style guide as sample content
    (the renderer clears the body, so the sample is only there to show the styles)."""
    _docx()
    from docx import Document
    doc = Document()
    ensure_styles(doc)
    _clear_body(doc)
    s = _setup_page(doc, "a4paper")
    from docx.shared import Mm
    s.top_margin, s.bottom_margin = Mm(15), Mm(14)
    s.left_margin = s.right_margin = Mm(17)
    s.header_distance, s.footer_distance = Mm(8), Mm(7)
    doc.add_paragraph("Job Hunt Kit - Word template", style="CV Name")
    doc.add_paragraph("Edit the styles, not this text", style="CV Title")
    doc.add_paragraph("This text is replaced on every build. Change the look by modifying the "
                      "styles below (Home > Styles > right-click > Modify), keep their names, "
                      "and save this file as templates/jobcv-template.docx in your workspace.",
                      style="CV Contact")
    for name in PARA_STYLES:
        if name in ("CV Name", "CV Title", "CV Contact"):
            continue
        doc.add_paragraph(f"{name}: sample text in this style", style=name)
    p = doc.add_paragraph(style="CV Body")
    for name in CHAR_STYLES:
        p.add_run(f"{name} (character style)", style=name)
        p.add_run("   ")
    doc.save(str(path))


def main() -> int:
    ap = argparse.ArgumentParser(description="Render content.json to LaTeX and/or Word.")
    ap.add_argument("folder", type=Path, nargs="?")
    ap.add_argument("--format", choices=("latex", "docx", "both"), default="docx")
    ap.add_argument("--templates", type=Path, help="workspace folder with jobcv.sty / jobcv-template.docx overrides")
    ap.add_argument("--make-template", type=Path, metavar="OUT.docx", help=argparse.SUPPRESS)
    args = ap.parse_args()
    try:
        if args.make_template:
            make_template(args.make_template)
            print(f"wrote {args.make_template}")
            return 0
        if not args.folder:
            ap.error("folder is required")
        for p in render(args.folder.resolve(), args.format, args.templates):
            print(f"wrote {p.name}")
    except ContentError as e:
        print(f"error: {e}", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
